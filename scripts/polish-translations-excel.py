#!/usr/bin/env python3
"""Apply header styling to Monlam-Playground-Translations-EN-BO.xlsx (run after export-translations-excel.mts)."""
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError as exc:
    raise SystemExit("Install openpyxl in docs/localization/.venv") from exc

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "localization" / "Monlam-Playground-Translations-EN-BO.xlsx"


def main() -> None:
    if not XLSX.is_file():
        raise SystemExit(f"Missing {XLSX}; run: npx tsx scripts/export-translations-excel.mts")

    wb = load_workbook(XLSX)
    ws = wb["Translations"]

    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, 4):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="top")
        for col in (2, 3):
            ws.cell(row=row, column=col).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 58
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"

    wb.save(XLSX)
    print(f"Polished {XLSX}")


if __name__ == "__main__":
    main()
